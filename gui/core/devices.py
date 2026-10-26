#######################################################################################################################
#######################################################################################################################
# Title:        Power Electronics Simulation Toolkit (PyPowerSim)
# Topic:        Power Electronics
# File:         devices.py
# Date:         26.10.2026
# Author:       Dr. Pascal A. Schirmer
# Version:      V.1.0
# Copyright:    Pascal Schirmer
#######################################################################################################################
#######################################################################################################################

#######################################################################################################################
# Function Description
#######################################################################################################################
"""
Reading and writing of the device parameter workbooks (para/Swi, para/Cap, para/Tra) for the input charts and the
editor. The cell positions are the ones of the loaders of the toolkit (src/data/loadParaSwi.py, loadParaCap.py,
loadParaTra.py):

- sheet 'electrical': parameter table (Parameter | Description | Model | Symbol | Typical | Value-1 … Value-10 | Unit)
  in the first lenElec rows; look-up tables (LUTs) of 10 rows x 10 temperature columns (B:K) in blocks of 20 rows,
  the rows belong to the vector of the x-axis parameter (e.g. If), the columns to the vector Tj
- sheet 'thermal': Foster networks (Rth_xx, Cth_xx), branches in Value-1 … Value-10, Typical = total
- sheet 'readme': author, date, datasheet and the ratings of the device

Saving: the workbook is opened with the cached values (formulas of the template, e.g. LUT axes '=F$16' or scaled
LUTs '=C55/1.42', become values), the edited cells are written and the result is checked with the loader of the
toolkit. openpyxl cannot recalculate formulas and the loaders read cached values, so keeping formulas would leave
the loaders with empty cells.
"""

#######################################################################################################################
# Import libs
#######################################################################################################################
from gui.core.files import FOLDERS, file_path

import contextlib
import copy
import io
import os

import numpy as np
import openpyxl
import pandas as pd

#######################################################################################################################
# Layout (Excel rows, 1-based) as in the loaders of the toolkit
#######################################################################################################################
COLS = ['Parameter', 'Description', 'Model', 'Symbol', 'Typical'] + [f"Value-{i}" for i in range(1, 11)] + ['Unit']
VALUES = [f"Value-{i}" for i in range(1, 11)]
N_LUT = 10                                                            # rows and temperature columns of a LUT
# kind -> rows of the parameter tables, first data row of every LUT and its x-axis vector (y-axis: Tj)
LAYOUT = {
    'swi': {'title': 'Switch', 'elec': 24, 'ther': 6,
            'luts': {'Vf': (35, 'If'), 'Eon': (55, 'If'), 'Eoff': (75, 'If'), 'Erec': (95, 'If'),
                     'Vfd': (115, 'Ifd'), 'Ciss': (135, 'Vf'), 'Coss': (155, 'Vf'), 'Crss': (175, 'Vf')}},
    'cap': {'title': 'Capacitor', 'elec': 10, 'ther': 4,
            'luts': {'C': (33, 'f'), 'tan': (53, 'f'), 'Kr': (73, 'f'), 'ESR': (93, 'f')}},
    'tra': {'title': 'Transformer', 'elec': 23, 'ther': 8,
            'luts': {'Rp': (35, 'f'), 'Rs': (55, 'f'), 'Lm': (75, 'B'), 'Rc1': (95, 'B'), 'Rc2': (115, 'Vp'),
                     'Cp': (135, 'f'), 'Cs': (155, 'f'), 'Cps': (175, 'f')}},
}
# unit of the LUT values (the block titles give the quantity; units as used by the loss models)
LUT_UNITS = {'Vf': 'V', 'Vfd': 'V', 'Eon': 'J', 'Eoff': 'J', 'Erec': 'J', 'Ciss': 'F', 'Coss': 'F', 'Crss': 'F',
             'C': 'F', 'tan': '-', 'Kr': '-', 'ESR': 'Ω', 'Rp': 'Ω', 'Rs': 'Ω', 'Lm': 'H', 'Rc1': 'Ω', 'Rc2': 'Ω',
             'Cp': 'F', 'Cs': 'F', 'Cps': 'F'}
# Foster networks of the thermal sheet: suffix -> meaning
THER_PATHS = {'JC': 'junction → case (transistor)', 'DC': 'junction → case (diode)', 'CA': 'case → ambient',
              'PC': 'primary winding → core', 'SC': 'secondary winding → core', 'CC': 'core → case'}


def kind_of(folder_key):
    """'swi' | 'cap' | 'tra' for the folder keys of the adapter."""
    return folder_key if folder_key in LAYOUT else None


def _num(v):
    try:
        x = float(v)
        return x if np.isfinite(x) else None
    except (TypeError, ValueError):
        return None


#######################################################################################################################
# Reading
#######################################################################################################################
def _grid(ws, rows, cols=16):
    out = [[None] * cols for _ in range(rows)]
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=rows, max_col=cols, values_only=True)):
        out[r][:len(row)] = row
    return out


def _table(grid, n):
    """Parameter table (Excel rows 2 … n+1) as DataFrame with the column names of the template."""
    rows = [list(grid[r][:len(COLS)]) + [None] * (len(COLS) - len(grid[r][:len(COLS)])) for r in range(1, n + 1)]
    df = pd.DataFrame(rows, columns=COLS)
    df.index = range(2, n + 2)                                        # Excel rows
    return df


def read(kind, name):
    """
    Device workbook as dict: 'readme' [(item, value, unit)], 'elec' / 'ther' DataFrames of the parameter tables,
    'luts' {symbol: {'title', 'xlabel', 'x', 'tj', 'z' (10x10, rows x, columns Tj), 'xsym', 'unit'}}.
    """
    L = LAYOUT[kind]
    wb = openpyxl.load_workbook(file_path(kind, name), data_only=True, read_only=True)
    try:
        el = _grid(wb['electrical'], 200)
        th = _grid(wb['thermal'], max(L['ther'] + 2, 12))
        rd = _grid(wb['readme'], 20, 3) if 'readme' in wb.sheetnames else [[None] * 3] * 20
    finally:
        wb.close()
    elec, ther = _table(el, L['elec']), _table(th, L['ther'])
    vec = {str(s): [_num(v) for v in row[VALUES].tolist()] for s, row in elec.set_index('Symbol', drop=False)
           .iterrows() if s is not None}
    luts = {}
    for sym, (r0, xsym) in L['luts'].items():
        x = [v for v in vec.get(xsym, []) if v is not None][:N_LUT]
        tj = [v for v in vec.get('Tj', []) if v is not None][:N_LUT]
        z = [[_num(el[r0 - 1 + i][1 + j]) for j in range(N_LUT)] for i in range(N_LUT)]
        luts[sym] = {'title': str(el[r0 - 4][0] or sym), 'xlabel': str(el[r0 - 3][0] or xsym), 'xsym': xsym,
                     'x': x, 'tj': tj, 'z': z, 'unit': LUT_UNITS.get(sym, ''), 'row': r0}
    readme = []
    for r in rd[7:]:                                                  # table 'Parameter | Value | Unit' from row 8
        if r[0] in (None, ''):
            break                                                     # below: version history of the workbook
        readme.append((str(r[0]), r[1], r[2]))
    info = {str(r[0]): r[1] for r in rd[:3] if r[0] not in (None, '')}
    return {'kind': kind, 'name': name, 'readme': readme, 'info': info, 'elec': elec, 'ther': ther, 'luts': luts}


def lut_frame(lut):
    """LUT as DataFrame: rows = x-axis values, columns = Tj values (only the defined part of the 10 x 10 block)."""
    nx, nt = len(lut['x']), len(lut['tj'])
    z = np.array([[np.nan if v is None else v for v in row[:nt]] for row in lut['z'][:nx]], dtype=float)
    df = pd.DataFrame(z.reshape(nx, nt), index=lut['x'], columns=[f"Tj = {t:g} °C" for t in lut['tj']])
    df.index.name = lut['xlabel']
    return df


def foster(dev):
    """Foster networks {path: (Rth [K/W] list, Cth [Ws/K] list)} of the thermal sheet."""
    out = {}
    ther = dev['ther'].set_index('Symbol')
    for sym in ther.index:
        if isinstance(sym, str) and sym.startswith('Rth_'):
            p = sym[4:]
            R = [_num(v) for v in ther.loc[sym, VALUES].tolist()]
            C = [_num(v) for v in ther.loc[f"Cth_{p}", VALUES].tolist()] if f"Cth_{p}" in ther.index else []
            pairs = [(r, c) for r, c in zip(R, C) if r is not None and c is not None]
            out[p] = ([r for r, _ in pairs], [c for _, c in pairs])
    return out


def zth(R, C, t):
    """Thermal impedance of a Foster network: Zth(t) = Σ Ri·(1 − exp(−t/(Ri·Ci)))."""
    t = np.asarray(t, dtype=float)[:, None]
    R, C = np.asarray(R, dtype=float), np.asarray(C, dtype=float)
    return np.sum(R * (1 - np.exp(-t / (R * C))), axis=1)


#######################################################################################################################
# Writing
#######################################################################################################################
def _clean(v):
    if v is None or (isinstance(v, float) and not np.isfinite(v)) or (isinstance(v, str) and v.strip() == ''):
        return None
    if isinstance(v, (np.floating, np.integer)):
        v = v.item()
    if isinstance(v, float) and v == int(v) and abs(v) < 1e9:
        return int(v)
    return v


def save(dev, dst, setup):
    """
    Writes the (edited) device dict into a copy of its workbook (all formulas replaced by their values) under
    para/<folder>/<dst>.xlsx and loads it with the loader of the toolkit; returns (path, para, log).
    """
    kind, L = dev['kind'], LAYOUT[dev['kind']]
    wb = openpyxl.load_workbook(file_path(kind, dev['name']), data_only=True)
    for sheet, key in (('electrical', 'elec'), ('thermal', 'ther')):
        ws = wb[sheet]
        for r, row in dev[key].iterrows():
            for c, col in enumerate(COLS, start=1):
                if col in row.index and col != 'Parameter':
                    ws.cell(int(r), c).value = _clean(row[col])
    ws = wb['electrical']
    vec = {str(s): [_num(v) for v in row[VALUES].tolist()] for s, row in dev['elec'].set_index('Symbol', drop=False)
           .iterrows() if s is not None}
    for sym, (r0, xsym) in L['luts'].items():
        lut = dev['luts'][sym]
        x, tj = [v for v in vec.get(xsym, []) if v is not None], [v for v in vec.get('Tj', []) if v is not None]
        for j in range(N_LUT):                                        # axes as values (template: formulas)
            ws.cell(r0 - 1, 2 + j).value = tj[j] if j < len(tj) else None
        for i in range(N_LUT):
            ws.cell(r0 + i, 1).value = x[i] if i < len(x) else None
            for j in range(N_LUT):
                ws.cell(r0 + i, 2 + j).value = _clean(lut['z'][i][j])
    folder = FOLDERS[kind]
    out = os.path.join(folder, f"{dst}.xlsx")
    wb.save(out)
    return (out,) + check(kind, dst, setup)


def check(kind, name, setup):
    """Loads a device workbook with the loader of the toolkit; returns (para or None, console output / error)."""
    from src.data.loadParaSwi import loadParaSwi
    from src.data.loadParaCap import loadParaCap
    from src.data.loadParaTra import loadParaTra
    from gui import adapter as A
    fnc = {'swi': loadParaSwi, 'cap': loadParaCap, 'tra': loadParaTra}[kind]
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            para = fnc(name, A.paths()['parPath'], copy.deepcopy(setup))
        return para, buf.getvalue()
    except Exception as e:
        return None, buf.getvalue() + f"\nERROR: {type(e).__name__}: {e}"
